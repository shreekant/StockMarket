import requests
from bs4 import BeautifulSoup
import re
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from stock_symbols import stock_symbols
import time
import os
import streamlit as st

def fetch_stock_details(url):
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
    }
    retries = 5
    for attempt in range(retries):
        try:
            response = requests.get(url, headers=headers, timeout=10)
            if response.status_code == 200:
                soup = BeautifulSoup(response.text, 'html.parser')

                pe_ratio_element = soup.find('span', string=lambda x: x and 'P/E' in x)
                high_low_element = soup.find('span', string=lambda x: x and 'High / Low' in x)
                roce_element = soup.find('span', string=lambda x: x and 'ROCE' in x)
                current_price_element = soup.find('span', string=lambda x: x and 'Current Price' in x)

                pe_ratio = pe_ratio_element.find_next_sibling('span').text.strip() if pe_ratio_element else None
                high_low_value = high_low_element.find_next_sibling('span').text.strip() if high_low_element else None
                roce_value = roce_element.find_next_sibling('span').text.strip() if roce_element else None
                current_price = current_price_element.find_next_sibling('span').text.strip() if current_price_element else None

                return pe_ratio, high_low_value, roce_value, current_price
            elif response.status_code == 429:
                wait_time = int(response.headers.get("Retry-After", (attempt + 1) * 5))
                print(f"Rate limit exceeded for {url}. Waiting {wait_time} seconds before retrying...")
                time.sleep(wait_time)
            else:
                print(f"Failed to fetch data from {url} with status code {response.status_code}")
                break
        except Exception as e:
            print(f"Error requesting {url}: {e}")
            time.sleep(2)
    return None, None, None, None

def get_details_of_stock(stock_symbol):
    urls = [
        f"https://www.screener.in/company/{stock_symbol}/consolidated",
        f"https://www.screener.in/company/{stock_symbol}"
    ]

    pe_ratio, high_low_value, roce_value, current_price = fetch_stock_details(urls[0])

    # Check if any key data points are missing, if so, try the second URL
    if not pe_ratio or not high_low_value or not roce_value or not current_price:
        pe_ratio, high_low_value, roce_value, current_price = fetch_stock_details(urls[1])

    if pe_ratio or high_low_value or roce_value or current_price:
        roce_value = re.sub(r'[^\d.]', '', roce_value) if roce_value else None  # Remove special characters except decimal point
        current_price = re.sub(r'[^\d.]', '', current_price) if current_price else None  # Remove special characters except decimal point

        # Split High/Low value
        if high_low_value:
            high, low = high_low_value.split('/')
            high = re.sub(r'[^\d.]', '', high.strip())
            low = re.sub(r'[^\d.]', '', low.strip())
        else:
            high, low = None, None

        # Convert high and low to float if possible
        try:
            high_value = float(high)
        except (ValueError, TypeError):
            high_value = None

        try:
            low_value = float(low)
        except (ValueError, TypeError):
            low_value = None

        # Convert current price to float if possible
        try:
            current_price_value = float(current_price)
        except (ValueError, TypeError):
            current_price_value = None

        # Calculate percentage change from high and low to current price
        if high_value is not None and current_price_value is not None:
            percentage_change_from_high = ((current_price_value - high_value) / high_value) * 100
        else:
            percentage_change_from_high = None

        if low_value is not None and current_price_value is not None:
            percentage_change_from_low = ((current_price_value - low_value) / low_value) * 100
        else:
            percentage_change_from_low = None

        # Convert PE ratio to float if possible
        try:
            pe_ratio_value = float(pe_ratio)
        except (ValueError, TypeError):
            pe_ratio_value = None

        return {
            "Symbol": stock_symbol,
            "PE Ratio": pe_ratio_value,
            "Current Price": current_price_value,
            "High": high_value,
            "Low": low_value,
            "ROCE": float(roce_value) if roce_value else None,
            "Percentage change from high": percentage_change_from_high,
            "Percentage change from low": percentage_change_from_low
        }
    return None

def generate_data():
    stock_details_list = []

    for symbol in stock_symbols:
        stock_details = get_details_of_stock(symbol)
        if stock_details:
            stock_details_list.append(stock_details)
        else:
            print(f"No data found for {symbol}")

        # Pace requests to respect screener.in rate limits
        time.sleep(2)

    # Create a DataFrame and save to Excel
    df = pd.DataFrame(stock_details_list)
    excel_filename = 'stock_details.xlsx'
    df.to_excel(excel_filename, index=False)

    # Load the workbook and select the active worksheet
    wb = load_workbook(excel_filename)
    ws = wb.active

    # Define the red fill for conditional formatting
    red_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
    orange_fill = PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid")
    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        pe_ratio_cell = row[1]  # Assuming PE Ratio is in the second column
        if pe_ratio_cell.value is not None:
            if pe_ratio_cell.value > 50:
                for cell in row:
                    cell.fill = red_fill
            elif 30 <= pe_ratio_cell.value <= 50:
                for cell in row:
                    cell.fill = orange_fill
            elif 24 <= pe_ratio_cell.value < 30:
                for cell in row:
                    cell.fill = yellow_fill

    # Auto-adjust column widths based on the title
    for column in ws.columns:
        max_length = 0
        column = list(column)
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(cell.value)
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column[0].column_letter].width = adjusted_width

    # Save the workbook with the applied formatting
    wb.save(excel_filename)

    print(f'Stock details saved to {excel_filename}')
    return excel_filename

def main():
    st.set_page_config(page_title="Stock Details Generator", page_icon="📈", layout="wide")
    st.title("📈 Stock Market Content Generator")
    st.write("Generate and download formatted stock market analysis reports.")

    col1, col2 = st.columns(2)

    with col1:
        st.subheader("Generate Content")
        generate_clicked = st.button("🚀 Generate Content", type="primary", use_container_width=True)

    with col2:
        st.subheader("Download Content")
        if generate_clicked:
            with st.spinner("Generating content..."):
                excel_filename = generate_data()
            st.success("Done! Content is ready to be downloaded.")
            with open(excel_filename, "rb") as f:
                st.download_button(
                    label="📥 Download Content (Excel)",
                    data=f.read(),
                    file_name="stock_details.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True
                )
        else:
            st.info("Please click 'Generate Content' first to prepare the download.")

if __name__ == "__main__":
    main()